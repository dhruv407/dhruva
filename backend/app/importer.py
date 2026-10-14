import hashlib
import io
import json
import os
import zipfile
from collections import Counter, defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

from .models import Snapshot, SourceRow
from .rules import LOCATIONS, MATERIALS

STOCK_REQUIRED = [
    "label_id",
    "branch",
    "material",
    "item",
    "karatage",
    "variety",
    "supplier",
    "received_at",
    "usable",
]
STOCK_OPTIONAL = [
    "set_id",
    "set_size",
    "weight_g",
    "carats",
    "metal_cost",
    "stone_cost",
    "other_cost",
    "landed_cost",
    "selling_price",
    "net_sale_value",
]
SALES_REQUIRED = [
    "sale_id",
    "label_id",
    "branch",
    "material",
    "item",
    "karatage",
    "variety",
    "sold_at",
    "selling_price",
]
SALES_OPTIONAL = ["buying_cost"]
MAX_BYTES = 15 * 1024 * 1024
MAX_ROWS = 30000


def workbook(content):
    if len(content) > MAX_BYTES:
        raise ValueError("Workbook exceeds 15 MB")
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as z:
            if sum(i.file_size for i in z.infolist()) > 100 * 1024 * 1024:
                raise ValueError("Expanded workbook exceeds 100 MB")
        return load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    except (zipfile.BadZipFile, KeyError, OSError) as e:
        raise ValueError("Provide a valid .xlsx workbook") from e


def inspect_workbook(content):
    wb = workbook(content)
    try:
        sheets = []
        for ws in wb:
            first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
            sheets.append(
                {
                    "name": ws.title,
                    "headers": [str(v).strip() if v is not None else "" for v in first],
                    "estimated_rows": max(0, (ws.max_row or 1) - 1),
                }
            )
        return {
            "sheets": sheets,
            "stock_required": STOCK_REQUIRED,
            "stock_optional": STOCK_OPTIONAL,
            "sales_required": SALES_REQUIRED,
            "sales_optional": SALES_OPTIONAL,
        }
    finally:
        wb.close()


def excel_rows(content, sheet):
    wb = workbook(content)
    try:
        if sheet not in wb.sheetnames:
            raise ValueError("Selected sheet does not exist")
        ws = wb[sheet]
        stream = ws.iter_rows()
        first = next(stream, ())
        headers = [str(c.value).strip() if c.value is not None else "" for c in first]
        nonempty = [h for h in headers if h]
        if not nonempty or len(nonempty) != len(set(nonempty)):
            raise ValueError("Headers must be nonempty and unique")
        rows = []
        for number, cells in enumerate(stream, 2):
            if number > MAX_ROWS + 1:
                raise ValueError("Workbook exceeds 30,000 rows per sheet")
            if not any(c.value is not None for c in cells):
                continue
            if any(c.data_type == "f" for c in cells):
                raise ValueError(
                    f"Row {number}: formula cells are not accepted; export values"
                )
            rows.append((number, {h: c.value for h, c in zip(headers, cells) if h}))
        return headers, rows
    finally:
        wb.close()


def iso_date(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return date.fromisoformat(str(value).strip()).isoformat()


def text(value):
    if value is None or not str(value).strip():
        raise ValueError("value is required")
    return str(value).strip()


def amount(value):
    if value is None or str(value).strip() == "":
        return None
    try:
        n = Decimal(str(value))
        if not n.is_finite() or n < 0 or n > Decimal("1000000000"):
            raise ValueError("amount must be finite, nonnegative and ≤ 1 billion")
        return str(n)
    except InvalidOperation as e:
        raise ValueError("invalid numeric amount") from e


def cost_paise(facts):
    if facts.get("landed_cost") is not None:
        return int((Decimal(facts["landed_cost"]) * 100).quantize(Decimal("1")))
    values = [facts.get(k) for k in ("metal_cost", "stone_cost", "other_cost")]
    if any(v is None for v in values):
        return None
    return int((sum(Decimal(v) for v in values) * 100).quantize(Decimal("1")))


def preview(db, content, filename, kind, sheet, mapping, as_of):
    as_of = iso_date(as_of)
    if date.fromisoformat(as_of) > date.today():
        raise ValueError("Snapshot date cannot be in the future")
    if kind not in ("stock", "sales"):
        raise ValueError("Import kind must be stock or sales")
    required = STOCK_REQUIRED if kind == "stock" else SALES_REQUIRED
    optional = STOCK_OPTIONAL if kind == "stock" else SALES_OPTIONAL
    headers, rows = excel_rows(content, sheet)
    if any(not mapping.get(k) for k in required):
        raise ValueError(f"Map all required fields: {', '.join(required)}")
    if any(
        k not in required + optional or v not in headers for k, v in mapping.items()
    ):
        raise ValueError("Mapping contains unknown fields or missing headers")
    if len(set(mapping.values())) != len(mapping):
        raise ValueError("A source column cannot map to multiple fields")
    errors, warnings, parsed = [], [], []
    seen = set()
    sales_labels = set()
    for number, raw in rows:
        facts = {}
        for field in required + optional:
            value = raw.get(mapping.get(field))
            try:
                if field not in mapping and field in optional:
                    facts[field] = None
                elif field in ("received_at", "sold_at"):
                    facts[field] = iso_date(value)
                    if facts[field] > as_of:
                        raise ValueError("date is after snapshot date")
                elif field in (
                    "weight_g",
                    "carats",
                    "metal_cost",
                    "stone_cost",
                    "other_cost",
                    "landed_cost",
                    "selling_price",
                    "net_sale_value",
                    "buying_cost",
                ):
                    facts[field] = amount(value)
                    if field in required and facts[field] is None:
                        raise ValueError("value is required")
                elif field == "set_size":
                    facts[field] = None if value is None else int(text(value))
                    if facts[field] is not None and facts[field] < 1:
                        raise ValueError("set_size must be a positive integer")
                elif field == "usable":
                    normalized = str(value).strip().lower()
                    if normalized not in ("true", "false", "yes", "no", "1", "0"):
                        raise ValueError("usable must explicitly be true/false")
                    facts[field] = normalized in ("true", "yes", "1")
                elif field == "set_id":
                    facts[field] = (
                        str(value).strip()
                        if value is not None and str(value).strip()
                        else None
                    )
                else:
                    if field in ("label_id", "sale_id") and not isinstance(value, str):
                        raise ValueError(
                            "identifier must be text, preserving leading zeros"
                        )
                    facts[field] = text(value)
                    if field == "branch" and facts[field] not in LOCATIONS:
                        raise ValueError(
                            "unknown branch; map explicitly to NV/LN/AR/RO/IK/HO/CZ"
                        )
                    if field == "material" and facts[field] not in MATERIALS:
                        raise ValueError(
                            "material must be gold, natural_diamond or lab_diamond"
                        )
            except (ValueError, TypeError) as e:
                errors.append({"row": number, "field": field, "message": str(e)})
        key = facts.get("label_id" if kind == "stock" else "sale_id")
        if key in seen:
            errors.append(
                {
                    "row": number,
                    "field": "identifier",
                    "message": "Duplicate identifier",
                }
            )
        seen.add(key)
        if kind == "sales":
            label = facts.get("label_id")
            if label in sales_labels:
                errors.append(
                    {
                        "row": number,
                        "field": "label_id",
                        "message": "Repeated sale of a label requires explicit return/re-sale reconciliation",
                    }
                )
            sales_labels.add(label)
        if kind == "stock" and cost_paise(facts) is None:
            warnings.append(
                {
                    "row": number,
                    "message": "Landed cost unknown; excluded from cost-balanced allocation",
                }
            )
        safe_raw = json.loads(
            json.dumps(
                raw,
                default=lambda v: (
                    v.isoformat() if isinstance(v, (date, datetime)) else str(v)
                ),
            )
        )
        parsed.append((number, safe_raw, facts))
    if kind == "stock":
        sets = defaultdict(list)
        for number, raw, f in parsed:
            if f.get("set_id"):
                sets[f["set_id"]].append((number, f))
        for key, members in sets.items():
            if (
                any(f.get("set_size") != len(members) for _, f in members)
                or len(
                    {
                        (f.get("branch"), f.get("material"), f.get("usable"))
                        for _, f in members
                    }
                )
                != 1
            ):
                errors.append(
                    {
                        "row": members[0][0],
                        "field": "set_id",
                        "message": f"Set {key}: incomplete set or mixed branch/material/usability",
                    }
                )
    digest = hashlib.sha256(content).hexdigest()
    root = Path(os.getenv("RG_DATA_DIR", "/workspace/dhruva/data")) / "sources"
    root.mkdir(parents=True, exist_ok=True)
    path = root / f"{digest}.xlsx"
    if not path.exists():
        path.write_bytes(content)
    snap = Snapshot(
        kind=kind,
        filename=Path(filename).name[:255],
        sha256=digest,
        blob_path=str(path),
        as_of=as_of,
        mapping={"sheet": sheet, "columns": mapping},
        errors=errors,
        warnings=warnings,
        coverage={},
    )
    db.add(snap)
    db.flush()
    db.add_all(
        SourceRow(snapshot_id=snap.id, row_number=n, raw=r, facts=f)
        for n, r, f in parsed
    )
    return snap


def reconcile(
    snap, rows, coverage, expected_counts, expected_costs, note, history_start=None
):
    if snap.status != "preview":
        raise ValueError("Snapshot has already been reconciled")
    if snap.errors:
        raise ValueError("Fix blocking import errors and import a new snapshot")
    if not note.strip():
        raise ValueError("Reconciliation note is required")
    if set(coverage) != set(LOCATIONS) or any(
        v not in ("missing", "complete", "confirmed_zero") for v in coverage.values()
    ):
        raise ValueError("Declare coverage for every location")
    counts = Counter(r.facts["branch"] for r in rows)
    totals = {}
    for location, status in coverage.items():
        if status == "confirmed_zero" and counts[location]:
            raise ValueError(f"{location}: confirmed zero contradicts imported rows")
        if status != "missing":
            if (
                location not in expected_counts
                or expected_counts[location] != counts[location]
            ):
                raise ValueError(
                    f"{location}: report label/row count does not reconcile"
                )
        if status == "missing" and counts[location]:
            raise ValueError(
                f"{location}: imported rows require complete coverage declaration"
            )
        costs = (
            [cost_paise(r.facts) for r in rows if r.facts["branch"] == location]
            if snap.kind == "stock"
            else []
        )
        totals[location] = None if any(c is None for c in costs) else sum(costs)
        if location in expected_costs:
            if (
                snap.kind != "stock"
                or totals[location] is None
                or abs(
                    totals[location] - int(Decimal(str(expected_costs[location])) * 100)
                )
                > 1
            ):
                raise ValueError(f"{location}: cost control total does not reconcile")
    if snap.kind == "sales":
        if not history_start:
            raise ValueError("Sales history requires an explicit coverage start date")
        history_start = iso_date(history_start)
        if history_start > snap.as_of or any(
            r.facts["sold_at"] < history_start for r in rows
        ):
            raise ValueError("Sales rows fall outside the declared history window")
    snap.coverage = coverage
    snap.reconciliation = {
        "counts": dict(counts),
        "cost_paise": totals,
        "expected_counts": expected_counts,
        "expected_costs": expected_costs,
        "note": note,
        "history_start": history_start,
    }
    snap.status = "reconciled"
